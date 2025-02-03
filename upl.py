# Nieuwe PowerPoint van 5 dia's met de gevraagde focus
from pptx import Presentation


prs = Presentation()

# Slide 1: Titel (Introductie)
slide_1 = prs.slides.add_slide(prs.slide_layouts[0])
title = slide_1.shapes.title
title.text = "Het Werk der Daklozen: Hulp voor kwetsbare mensen"
subtitle = slide_1.placeholders[1]
subtitle.text = "Hoe de organisatie te werk gaat en mensen helpt"

# Slide 2: Over Het Werk der Daklozen
slide_2 = prs.slides.add_slide(prs.slide_layouts[1])
title_2 = slide_2.shapes.title
title_2.text = "Over Het Werk der Daklozen"
content_2 = slide_2.shapes.placeholders[1].text_frame
content_2.text = "Belgische organisatie die sinds 1908 hulp biedt aan daklozen en kwetsbare groepen."

# Slide 3: Hoe ze te werk gaan
slide_3 = prs.slides.add_slide(prs.slide_layouts[1])
title_3 = slide_3.shapes.title
title_3.text = "Hoe ze te werk gaan"
content_3 = slide_3.shapes.placeholders[1].text_frame
content_3.text = (
    "1. Begeleiding: hulp bij administratie, juridische zaken, huisvesting.\n"
    "2. Noodhulp: verstrekking van voedsel en materiële steun.\n"
    "3. Lange termijn: re-integratie via werk en opleiding."
)

# Slide 4: Wat ze doen voor mensen
slide_4 = prs.slides.add_slide(prs.slide_layouts[1])
title_4 = slide_4.shapes.title
title_4.text = "Wat ze doen voor mensen"
content_4 = slide_4.shapes.placeholders[1].text_frame
content_4.text = (
    "Ze bieden essentiële hulp, zoals voedsel, onderdak en begeleiding, en helpen mensen om weer zelfstandig te worden door werk en opleiding."
)

# Slide 5: Hoe mensen worden geholpen
slide_5 = prs.slides.add_slide(prs.slide_layouts[1])
title_5 = slide_5.shapes.title
title_5.text = "Hoe mensen worden geholpen"
content_5 = slide_5.shapes.placeholders[1].text_frame
content_5.text = (
    "Via hun diensten krijgen mensen ondersteuning om hun situatie te verbeteren. De organisatie richt zich zowel op noodhulp als op duurzame oplossingen."
)

# Opslaan van de presentatie
pptx_file_5slides = "/mnt/data/Het_Werk_der_Daklozen_Presentatie_5slides.pptx"
prs.save(pptx_file_5slides)

pptx_file_5slides
